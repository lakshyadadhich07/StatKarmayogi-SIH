import logging
import os
import re
from typing import List
from langchain_core.documents import Document as LCDocument
from langchain_community.document_loaders import PyPDFLoader
from pptx import Presentation

logger = logging.getLogger(__name__)


class DocumentParsingError(Exception):
    """Raised when parsing fails due to corrupt, invalid, or unsupported files."""
    pass


class EmptyDocumentError(DocumentParsingError):
    """Raised when no extractable text could be found in the document."""
    pass


def clean_text(text: str) -> str:
    """Conservatively cleans extracted text by normalizing line breaks and whitespace."""
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    joined = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", joined).strip()


def load_pdf_with_langchain(file_path: str) -> List[LCDocument]:
    """Loads PDF pages using LangChain's official PyPDFLoader."""
    if not os.path.exists(file_path):
        raise DocumentParsingError(f"File not found: {file_path}")

    try:
        loader = PyPDFLoader(file_path)
        raw_docs = loader.load()
    except Exception as e:
        logger.error(f"PyPDFLoader error for {file_path}: {e}")
        raise DocumentParsingError(f"Failed to load PDF via PyPDFLoader: {str(e)}")

    normalized_docs: List[LCDocument] = []
    for doc in raw_docs:
        cleaned = clean_text(doc.page_content)
        if cleaned:
            # LangChain PyPDFLoader provides 0-indexed 'page'; normalize to 1-indexed
            raw_page = doc.metadata.get("page", 0)
            page_num = int(raw_page) + 1
            meta = dict(doc.metadata)
            meta["page"] = page_num
            meta["page_number"] = page_num
            meta["source"] = os.path.basename(file_path)
            normalized_docs.append(LCDocument(page_content=cleaned, metadata=meta))

    return normalized_docs


def load_pptx_as_langchain_docs(file_path: str) -> List[LCDocument]:
    """Uses python-pptx solely as a non-AI format parser to produce standardized LangChain Document objects."""
    if not os.path.exists(file_path):
        raise DocumentParsingError(f"File not found: {file_path}")

    try:
        prs = Presentation(file_path)
    except Exception as e:
        logger.error(f"python-pptx error for {file_path}: {e}")
        raise DocumentParsingError(f"Failed to parse PowerPoint presentation: {str(e)}")

    docs: List[LCDocument] = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_parts: List[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    p_text = "".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text
                    if p_text and p_text.strip():
                        slide_parts.append(p_text.strip())
            elif shape.has_table:
                for row in shape.table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                    if row_texts:
                        slide_parts.append(" | ".join(row_texts))

        slide_content = clean_text("\n".join(slide_parts))
        if slide_content:
            meta = {
                "source": os.path.basename(file_path),
                "page": slide_idx,
                "page_number": slide_idx,
                "slide_number": slide_idx,
            }
            docs.append(LCDocument(page_content=slide_content, metadata=meta))

    return docs


def load_document_pages(file_path: str, file_type: str) -> List[LCDocument]:
    """Loads document using PyPDFLoader for PDF or non-AI format parser for PPTX, returning LangChain Documents."""
    normalized_type = file_type.lower().strip().lstrip(".")

    if normalized_type == "pdf":
        docs = load_pdf_with_langchain(file_path)
    elif normalized_type == "pptx":
        docs = load_pptx_as_langchain_docs(file_path)
    elif normalized_type == "ppt":
        # Report explicit limitation for legacy binary OLE PPT format
        raise DocumentParsingError(
            "Legacy binary .ppt format is not directly supported; please convert to modern .pptx or .pdf format."
        )
    else:
        raise DocumentParsingError(
            f"Unsupported file type: {file_type}. Supported types: PDF, PPTX."
        )

    if not docs or not any(d.page_content.strip() for d in docs):
        logger.warning(f"No extractable text found in {file_path}")
        raise EmptyDocumentError("No extractable text found in document.")

    return docs
